from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def set_columns(section, count: int) -> None:
    section_properties = section._sectPr
    columns = section_properties.find(qn("w:cols"))
    if columns is None:
        columns = OxmlElement("w:cols")
        section_properties.append(columns)
    columns.set(qn("w:num"), str(count))
    columns.set(qn("w:space"), "360")


def set_run_font(run, size: float, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def format_document(input_path: Path, output_path: Path) -> None:
    document = Document(input_path)
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        set_columns(section, 2)

    for paragraph in document.paragraphs:
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.0
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            set_run_font(run, 9)

    if document.paragraphs:
        title = document.paragraphs[0]
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(8)
        for run in title.runs:
            set_run_font(run, 16, bold=True)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        set_run_font(run, 7.5)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: format_ieee_docx.py INPUT.docx OUTPUT.docx")
    format_document(Path(sys.argv[1]), Path(sys.argv[2]))